"""Export the letters as PDFs (all 250) and DOCX (a sample of 10), then check every file reads back.

Each PDF is laid out like a GP referral letter: the (fictional) practice's name, address and code, the date and
the e-RS reference, then the letter text, on A4 in Helvetica. No NHS logo or real branding, and a footer on every
page: "Synthetic test letter – not a real patient". Text is drawn as text (not images), so pypdf can read it.

The check reads each file back with kneedss.extract.read_letter and confirms every answer-key quote is still found
(spacing ignored, as in the app's quote check). No API calls.

    python scripts/export_pdfs.py
      -> data/output/letters_pdf/, data/output/test/letters_pdf/, data/output/letters_docx/
         data/output/letters_pdf.zip, letters_test_pdf.zip, letters_docx_sample.zip
"""
from __future__ import annotations

import io
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from kneedss import standards as S                       # noqa: E402
from kneedss.extract import quote_in_letter, read_letter  # noqa: E402

OUT = ROOT / "data" / "output"
SETS = {"dev": OUT, "test": OUT / "test"}
FOOTER = "Synthetic test letter – not a real patient"
DOCX_SAMPLE = ["L0001", "L0002", "L0003", "L0004", "L0005", "T0001", "T0002", "T0003", "T0004", "T0005"]

PAGE_W, PAGE_H = A4
MARGIN = 20 * mm
FONT, BOLD, SIZE, LEADING = "Helvetica", "Helvetica-Bold", 9.5, 13


def letter_meta(text: str) -> dict:
    """Practice, date and e-RS reference, read from the letter's own record block."""
    flat = re.sub(r"\s+", " ", text)
    code = re.search(r"\((Y99\d{3})\)", flat)
    date = re.search(r"Referral date: (\d{2}/\d{2}/\d{4})", flat)
    ubrn = re.search(r"UBRN: (\d{4} \d{4} \d{4})", flat)
    code = code.group(1) if code else S.PRACTICES[0][0]
    return {"code": code, "practice": dict(S.PRACTICES)[code], "address": S.PRACTICE_ADDRESSES[code],
            "date": date.group(1) if date else "", "ubrn": ubrn.group(1) if ubrn else ""}


def letter_pdf(text: str, meta: dict) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, pageCompression=1)
    c.setTitle("Synthetic referral letter")
    c.setAuthor(meta["practice"])
    width = PAGE_W - 2 * MARGIN
    page = 1

    def footer():
        c.setFont(FONT, 8)
        c.setFillGray(0.35)
        c.drawString(MARGIN, 12 * mm, FOOTER)
        c.drawRightString(PAGE_W - MARGIN, 12 * mm, f"Page {page}")
        c.setFillGray(0)

    def letterhead() -> float:
        y = PAGE_H - MARGIN
        c.setFont(BOLD, 13)
        c.drawString(MARGIN, y, meta["practice"])
        c.setFont(FONT, SIZE)
        c.drawString(MARGIN, y - 15, meta["address"])
        c.drawString(MARGIN, y - 28, f"Practice code {meta['code']}")
        c.drawRightString(PAGE_W - MARGIN, y, f"Date: {meta['date']}")
        c.drawRightString(PAGE_W - MARGIN, y - 15, f"e-RS UBRN: {meta['ubrn']}")
        c.setLineWidth(0.5)
        c.line(MARGIN, y - 38, PAGE_W - MARGIN, y - 38)
        return y - 56

    y = letterhead()
    for block in text.replace("\r", "").split("\x0c"):          # a form feed in a scanned-style letter = new page
        if block is not text and y < PAGE_H - MARGIN - 60:
            footer()
            c.showPage()
            page += 1
            y = PAGE_H - MARGIN
        for raw in block.split("\n"):
            line = re.sub(r" {2,}", "   ", raw.rstrip())            # keep columns readable, not huge gaps
            pieces = simpleSplit(line, FONT, SIZE, width) or [""]
            for piece in pieces:
                if y < MARGIN + 12 * mm:
                    footer()
                    c.showPage()
                    page += 1
                    y = PAGE_H - MARGIN
                c.setFont(FONT, SIZE)
                c.drawString(MARGIN, y, piece)
                y -= LEADING
    footer()
    c.save()
    return buf.getvalue()


def letter_docx(text: str, meta: dict) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(10)
    head = doc.add_paragraph()
    head.add_run(meta["practice"]).bold = True
    doc.add_paragraph(meta["address"])
    doc.add_paragraph(f"Practice code {meta['code']}    Date: {meta['date']}    e-RS UBRN: {meta['ubrn']}")
    for line in text.replace("\x0c", "").split("\n"):
        doc.add_paragraph(line)
    for section in doc.sections:
        section.footer.paragraphs[0].text = FOOTER
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def check(name: str, data: bytes, key: dict) -> dict:
    """Read the file back and look for every answer-key quote."""
    text = read_letter(name, data)
    missing = [f for f, q in key["evidence"].items() if not quote_in_letter(q, text)]
    return {"file": name, "quotes": len(key["evidence"]), "found": len(key["evidence"]) - len(missing),
            "missing": missing}


def main() -> None:
    results = []
    docx_dir = OUT / "letters_docx"
    docx_dir.mkdir(parents=True, exist_ok=True)
    for set_name, base in SETS.items():
        keys = {k["letter_id"]: k for k in map(json.loads, open(base / "answer_key.jsonl"))}
        pdf_dir = base / "letters_pdf"
        pdf_dir.mkdir(parents=True, exist_ok=True)
        for lid, key in keys.items():
            text = (base / "letters" / f"{lid}.txt").read_text()
            meta = letter_meta(text)
            pdf = letter_pdf(text, meta)
            (pdf_dir / f"{lid}.pdf").write_bytes(pdf)
            results.append({"set": set_name, "type": "pdf", **check(f"{lid}.pdf", pdf, key)})
            if lid in DOCX_SAMPLE:
                docx = letter_docx(text, meta)
                (docx_dir / f"{lid}.docx").write_bytes(docx)
                results.append({"set": set_name, "type": "docx", **check(f"{lid}.docx", docx, key)})

    for zip_name, folder, pattern in [("letters_pdf.zip", OUT / "letters_pdf", "*.pdf"),
                                      ("letters_test_pdf.zip", OUT / "test" / "letters_pdf", "*.pdf"),
                                      ("letters_docx_sample.zip", docx_dir, "*.docx")]:
        with zipfile.ZipFile(OUT / zip_name, "w", zipfile.ZIP_DEFLATED) as z:
            for f in sorted(folder.glob(pattern)):
                z.write(f, f.name)

    print("Read-back check (every answer-key quote found after reading the file):")
    for kind in ("pdf", "docx"):
        for set_name in SETS:
            rs = [r for r in results if r["type"] == kind and r["set"] == set_name]
            if not rs:
                continue
            files_ok = sum(not r["missing"] for r in rs)
            quotes, found = sum(r["quotes"] for r in rs), sum(r["found"] for r in rs)
            print(f"  {kind.upper():4} {set_name:4}: {files_ok}/{len(rs)} files pass ({files_ok / len(rs):.1%}); "
                  f"quotes {found}/{quotes} ({found / quotes:.1%})")
    failed = [r for r in results if r["missing"]]
    for r in failed[:20]:
        print(f"  MISSING {r['file']}: {', '.join(r['missing'])}")
    print(f"Zips: {', '.join(z for z in ['letters_pdf.zip', 'letters_test_pdf.zip', 'letters_docx_sample.zip'])} "
          f"in {OUT.relative_to(ROOT)}/")
    if failed:
        sys.exit(1)


if __name__ == "__main__":
    main()
