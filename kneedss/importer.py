"""Import patients from a spreadsheet (CSV or Excel). No LLM, no API calls.

1. read_table: the header row and data rows of a .csv or .xlsx file, all as text.
2. detect: match each column header to a field, by exact name, a known alias, or fuzzy matching (difflib, 0.8).
3. clean_row: validate one row with the same rules as kneedss/clean.py (NHS number mod-11 check, OKS 0-48,
   KL 0-4, BMI 12-80, dates, yes/no and sex labels). An invalid row is skipped with a reason; nothing is guessed.
   A missing optional value stays empty (the rules then send the patient to "Needs info").
4. plan: every row cleaned and classified with rules.decide() and rules.priority_score(). The backend matches
   on NHS number to update or create patients (backend/services.py: import_commit).
"""
from __future__ import annotations

import csv
import difflib
import io
import re
from datetime import date, datetime
from pathlib import Path

from kneedss import clean as C
from kneedss.rules import PATHWAYS, decide, priority_breakdown, priority_score
from kneedss.standards import is_valid_nhs_number

MAX_ROWS = 5000

# field -> (label, required, aliases). Names match the cleaned waiting list where it has the field.
FIELDS: dict[str, tuple[str, bool, list[str]]] = {
    "nhs_number": ("NHS number", True, ["nhs no", "nhs number", "nhsnumber", "nhs", "nhs num", "nhs id"]),
    "name": ("Name", True, ["patient name", "full name", "patient", "name"]),
    "date_of_birth": ("Date of birth", True, ["dob", "date of birth", "birth date", "birthdate", "born"]),
    "sex": ("Sex", False, ["gender", "sex", "sex code"]),
    "side": ("Knee", False, ["knee", "side", "laterality", "affected knee", "which knee"]),
    "oks": ("Oxford Knee Score", False, ["oks", "oxford knee score", "oxford score", "knee score", "oks now", "current oks"]),
    "oks_before": ("Earlier OKS", False, ["oks before", "previous oks", "earlier oks", "oks previous", "prior oks", "last oks"]),
    "kl_grade": ("KL grade", False, ["kl", "kl grade", "kellgren lawrence", "kellgren lawrence grade", "xray grade",
                                     "x ray grade", "xr grade", "kellgren"]),
    "bmi": ("BMI", False, ["bmi", "body mass index", "bmi kg m2", "bmi kgm2"]),
    "physio_weeks": ("Physio weeks", False, ["physio", "physio weeks", "physiotherapy", "physiotherapy weeks",
                                             "weeks of physio", "physio wks"]),
    "injections": ("Injections", False, ["injections", "steroid injections", "injection count", "inj"]),
    "night_pain": ("Night pain", False, ["night pain", "pain at night", "night"]),
    "lives_alone_or_carer": ("Lives alone or carer", False, ["lives alone or carer", "lives alone carer", "lives alone",
                                                             "carer", "social"]),
    "walking_aid": ("Walking aid", False, ["walking aid", "stick", "uses stick", "mobility aid", "walking stick"]),
    "diabetes": ("Diabetes", False, ["diabetes", "diabetic", "t2dm", "dm"]),
    "smoking_status": ("Smoking", False, ["smoking", "smoking status", "smoker", "smokes"]),
    "referral_date": ("Referral date", True, ["referral date", "referred", "referred on", "date referred", "referral"]),
    "practice_code": ("Practice code", False, ["practice", "practice code", "gp practice", "ods code", "practice ods"]),
}
REQUIRED = [f for f, (_, req, _) in FIELDS.items() if req]
YES = {"yes", "y", "true", "1", "x"}
NO = {"no", "n", "false", "0"}
NUMBER_RANGES = {"oks": C.RANGES["oks"], "oks_before": C.RANGES["oks_before"], "kl_grade": C.RANGES["kl_grade"],
                 "bmi": C.RANGES["bmi"], "physio_weeks": C.RANGES["physio_weeks"], "injections": (0, 20)}
INTEGER_FIELDS = {"oks", "oks_before", "kl_grade", "physio_weeks", "injections"}


def norm(header: str) -> str:
    """Lowercase, letters and digits only: 'NHS No.' -> 'nhsno'."""
    return re.sub(r"[^a-z0-9]", "", str(header).lower())


# ---------------------------------------------------------------------------- reading

def read_table(name: str, data: bytes) -> tuple[list[str], list[list[str]]]:
    """Headers and rows (as text) of a .csv or .xlsx file. Raises ValueError if it can't be read."""
    suffix = Path(name).suffix.lower()
    if suffix == ".csv":
        text = data.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        cell = lambda v: "" if v is None else v.date().isoformat() if isinstance(v, datetime) else \
            v.isoformat() if isinstance(v, date) else str(v)
        rows = [[cell(v) for v in r] for r in sheet.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
    else:
        raise ValueError("Only .csv or .xlsx")
    if not rows:
        raise ValueError("Empty file")
    headers = [h.strip() for h in rows[0]]
    body = [(r + [""] * len(headers))[: len(headers)] for r in rows[1:]]
    if len(body) > MAX_ROWS:
        raise ValueError(f"At most {MAX_ROWS:,} rows")
    return headers, body


# ---------------------------------------------------------------------------- detecting columns

def detect(headers: list[str]) -> list[dict]:
    """For each header: the field it matches and how sure the match is (exact / alias / fuzzy / none).
    Each field goes to one column at most: exact beats alias beats fuzzy, then the first column wins."""
    names = {norm(f): f for f in FIELDS}
    aliases = {norm(a): f for f, (_, _, al) in FIELDS.items() for a in al}
    candidates = list(names) + list(aliases)
    found = []
    for i, h in enumerate(headers):
        n = norm(h)
        if n in names:
            found.append((0, i, names[n], "exact"))
        elif n in aliases:
            found.append((1, i, aliases[n], "alias"))
        elif n and (close := difflib.get_close_matches(n, candidates, n=1, cutoff=0.8)):
            found.append((2, i, names.get(close[0]) or aliases[close[0]], "fuzzy"))
    out = [{"header": h, "field": None, "confidence": "none"} for h in headers]
    taken = set()
    for _, i, field, how in sorted(found):
        if field not in taken:
            taken.add(field)
            out[i] = {"header": headers[i], "field": field, "confidence": how}
    return out


def mapping_problems(mapping: dict[str, str | None]) -> list[str]:
    """Why a mapping can't be imported: required fields with no column, or a field given two columns."""
    used = [f for f in mapping.values() if f]
    problems = [f"Missing: {FIELDS[f][0]}" for f in REQUIRED if f not in used]
    problems += [f"Two columns for {FIELDS[f][0]}" for f in sorted({f for f in used if used.count(f) > 1})]
    problems += [f"Unknown field: {f}" for f in sorted(set(used) - set(FIELDS))]
    return problems


# ---------------------------------------------------------------------------- one row

def _yes_no(v: str, label: str):
    s = str(v).strip().lower()
    if s == "":
        return None
    if s in YES:
        return True
    if s in NO:
        return False
    raise ValueError(f"{label} '{v}' is not yes or no")


def clean_row(raw: dict[str, str], as_of: date = C.EXTRACT_DATE) -> dict:
    """One row (field -> text) to a clean record, or raise ValueError with the reason it is skipped."""
    get = lambda f: str(raw.get(f, "") or "").strip()
    rec: dict = {}
    for f in REQUIRED:
        if not get(f):
            raise ValueError(f"{FIELDS[f][0]} missing")

    nhs = re.sub(r"\D", "", get("nhs_number"))
    if not is_valid_nhs_number(nhs):
        raise ValueError(f"NHS number '{get('nhs_number')}' fails the check")
    rec["nhs_number"] = f"{nhs[:3]} {nhs[3:6]} {nhs[6:]}"
    rec["name"] = re.sub(r"\s+", " ", get("name"))

    for f in ("date_of_birth", "referral_date"):
        d = C._date(get(f))
        if d is None:
            raise ValueError(f"{FIELDS[f][0]} '{get(f)}' is not a date (dd/mm/yyyy or yyyy-mm-dd)")
        if date.fromisoformat(d) > as_of:
            raise ValueError(f"{FIELDS[f][0]} {d} is in the future")
        rec[f] = d
    rec["age"] = C.age_on(rec["date_of_birth"], as_of)
    if not C.RANGES["age"][0] <= rec["age"] <= C.RANGES["age"][1]:
        raise ValueError(f"Age {rec['age']} from the date of birth is not plausible")
    rec["weeks_waiting"] = (as_of - date.fromisoformat(rec["referral_date"])).days // 7

    for f, lookup, label in (("sex", C.SEX, "Sex"), ("side", C.LATERALITY, "Knee"), ("smoking_status", C.SMOKING, "Smoking")):
        v = get(f)
        if v and lookup.get(v.lower()) is None:
            raise ValueError(f"{label} '{v}' is not recognised")
        rec[f] = lookup.get(v.lower()) if v else None

    for f, (lo, hi) in NUMBER_RANGES.items():
        v = get(f)
        if not v:
            rec[f] = None
            continue
        n = C._num(v)
        if n is None:
            raise ValueError(f"{FIELDS[f][0]} '{v}' is not a number")
        if not lo <= n <= hi:
            raise ValueError(f"{FIELDS[f][0]} {v:s} is outside {lo}-{hi}")
        if f in INTEGER_FIELDS and n != int(n):
            raise ValueError(f"{FIELDS[f][0]} {v:s} must be a whole number")
        rec[f] = int(n) if f in INTEGER_FIELDS else round(n, 1)

    for f in ("night_pain", "lives_alone_or_carer", "walking_aid", "diabetes"):
        rec[f] = _yes_no(get(f), FIELDS[f][0])

    code = get("practice_code").upper()
    if code and not re.fullmatch(r"[A-Z]\d{5}", code):
        raise ValueError(f"Practice code '{get('practice_code')}' is not an ODS code (e.g. Y99001)")
    rec["practice_code"] = code or None
    return rec


def classify(rec: dict) -> dict:
    """The rules on a clean record: pathway, rule, reasons, weight support, priority and its parts."""
    f = {"oks": rec["oks"], "oks_before": rec["oks_before"], "kl": rec["kl_grade"], "physio_weeks": rec["physio_weeks"],
         "bmi": rec["bmi"], "weeks_waiting": rec["weeks_waiting"], "night_pain": rec["night_pain"],
         "lives_alone_or_carer": rec["lives_alone_or_carer"]}
    d = decide(f)
    return {**d, "label": PATHWAYS[d["pathway"]], "priority": priority_score(f), "priority_parts": priority_breakdown(f),
            "facts": {**f, "injections": rec["injections"], "walking_aid": rec["walking_aid"], "diabetes": rec["diabetes"],
                      "smoking": rec["smoking_status"], "side": rec["side"], "sex": rec["sex"], "age": rec["age"]}}


# ---------------------------------------------------------------------------- whole file

def plan(headers: list[str], rows: list[list[str]], mapping: dict[str, str | None]) -> dict:
    """Clean and classify every row with a confirmed mapping (header -> field or None). Rows are numbered as in
    the spreadsheet (header = row 1). A repeated NHS number later in the file is skipped."""
    cols = {i: mapping.get(h) for i, h in enumerate(headers) if mapping.get(h)}
    valid, skipped, seen = [], [], set()
    for n, row in enumerate(rows, start=2):
        raw = {field: row[i] for i, field in cols.items()}
        try:
            rec = clean_row(raw)
        except ValueError as e:
            skipped.append({"row": n, "nhs_number": raw.get("nhs_number", ""), "name": raw.get("name", ""), "reason": str(e)})
            continue
        if rec["nhs_number"] in seen:
            skipped.append({"row": n, "nhs_number": rec["nhs_number"], "name": rec["name"], "reason": "NHS number repeated in this file"})
            continue
        seen.add(rec["nhs_number"])
        valid.append({"row": n, "record": rec, "decision": classify(rec)})
    return {"valid": valid, "skipped": skipped}


def preview(name: str, data: bytes, mapping: dict[str, str | None] | None = None) -> dict:
    """What the clinician sees before importing: the detected (or given) mapping, the first 10 rows, problems,
    and how many rows would import."""
    headers, rows = read_table(name, data)
    columns = detect(headers)
    if mapping is not None:
        columns = [{**c, "field": mapping.get(c["header"]), "confidence": c["confidence"] if mapping.get(c["header"]) == c["field"] else "manual"}
                   for c in columns]
    chosen = {c["header"]: c["field"] for c in columns}
    problems = mapping_problems(chosen)
    result = plan(headers, rows, chosen) if not problems else {"valid": [], "skipped": []}
    return {"file": name, "columns": columns, "rows": rows[:10], "total_rows": len(rows), "problems": problems,
            "valid": len(result["valid"]), "skipped": result["skipped"],
            "fields": [{"field": f, "label": lab, "required": req} for f, (lab, req, _) in FIELDS.items()]}


def template_csv() -> str:
    """The import template: the defined headers and two example rows (synthetic, NHS test numbers)."""
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(list(FIELDS))
    w.writerow(["999 000 0018", "Jane Example", "1956-03-14", "Female", "Right", "22", "26", "3", "31.4", "8", "1",
                "yes", "no", "no", "no", "Never smoked", "2026-05-02", "Y99001"])
    w.writerow(["999 000 0026", "John Sample", "1949-11-02", "Male", "Left", "35", "", "2", "36.2", "4", "0",
                "no", "yes", "yes", "yes", "Ex-smoker", "2026-08-19", "Y99004"])
    return out.getvalue()
